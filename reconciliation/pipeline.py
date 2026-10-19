"""Input adapters, validation, and reconciliation. No source data is modified."""

import csv
import json
import os
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path


def decimal(value):
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("amount must be a plain decimal number") from None
    if not result.is_finite():
        raise ValueError("amount must be finite")
    return result


def read_source(spec, base):
    """Return headers and rows. Database credentials come from the environment."""
    kind = spec.get("type", "csv")
    if kind == "csv":
        with (base / spec["path"]).open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.DictReader(stream)
            headers = reader.fieldnames or []
            rows = list(reader)
    elif kind == "excel":
        from openpyxl import load_workbook

        workbook = load_workbook(base / spec["path"], read_only=True, data_only=True)
        try:
            sheet = workbook[spec["sheet"]] if "sheet" in spec else workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = list(next(iterator, ()))
            rows = [dict(zip(headers, row)) for row in iterator if any(v is not None for v in row)]
        finally:
            workbook.close()
    elif kind in ("sql", "oracle", "mysql"):
        from sqlalchemy import create_engine, text

        if kind in ("oracle", "mysql"):
            from sqlalchemy import URL

            database_field = "service" if kind == "oracle" else "database"
            fields = ("user", "password", "host", "port", database_field)
            names = {k: spec.get(f"{k}_env", f"{kind.upper()}_{k.upper()}") for k in fields}
            values = {k: os.environ.get(name) for k, name in names.items()}
            missing = [names[k] for k in fields if k != "port" and not values[k]]
            if missing:
                raise ValueError("missing database environment variables: " + ", ".join(missing))
            url = URL.create(
                "oracle+oracledb" if kind == "oracle" else "mysql+pymysql",
                username=values["user"], password=values["password"], host=values["host"],
                port=int(values["port"] or ("1521" if kind == "oracle" else "3306")),
                database=values[database_field] if kind == "mysql" else None,
                query={"service_name": values[database_field]} if kind == "oracle" else {},
            )
        else:
            url = os.environ.get(spec.get("url_env", "WAREHOUSE_URL"))
        if not url:
            raise ValueError("database URL environment variable is missing")
        # The database account must enforce read-only access; arbitrary SQL is not sandboxed.
        engine = create_engine(url, hide_parameters=True, connect_args=spec.get("connect_args", {}))
        try:
            with engine.connect() as connection:
                result = connection.execute(text(spec["query"]), spec.get("parameters", {}))
                headers = list(result.keys())
                rows = [dict(row) for row in result.mappings()]
        finally:
            engine.dispose()
    else:
        raise ValueError(f"unsupported source type: {kind}")
    if not headers or any(not isinstance(h, str) or not h for h in headers):
        raise ValueError("source must have nonempty text column headers")
    if len(set(headers)) != len(headers):
        raise ValueError("source has duplicate column headers")
    return headers, rows


def prepare(headers, rows, spec, keys, amounts, side):
    mapping = spec.get("columns", {})
    required = keys + amounts
    missing = [name for name in required if mapping.get(name, name) not in headers]
    if missing:
        raise ValueError(f"{side}: missing columns: {', '.join(missing)}")
    indexed, issues = {}, []
    counts = Counter()
    for number, row in enumerate(rows, start=2):
        try:
            if None in row or any(v is None for v in row.values()):
                raise ValueError("row contains missing cells or unexpected extra cells")
            key = tuple(str(row[mapping.get(k, k)]).strip() for k in keys)
            if any(not value for value in key):
                raise ValueError("matching key is blank")
            values = {a: decimal(row[mapping.get(a, a)]) for a in amounts}
            counts[key] += 1
            indexed.setdefault(key, values)
        except ValueError as exc:
            issues.append({"source": side, "row": number, "issue": str(exc)})
    duplicates = {key for key, count in counts.items() if count > 1}
    for key in sorted(duplicates):
        issues.append({"source": side, "row": "", "issue": f"duplicate key {key!r} ({counts[key]} records)"})
    return indexed, duplicates, issues


def reconcile(config, base):
    keys, amounts = config["keys"], config["amounts"]
    if (not keys or not amounts or any(not isinstance(x, str) or not x for x in keys + amounts)
            or len(set(keys + amounts)) != len(keys + amounts)):
        raise ValueError("keys and amounts must be nonempty, distinct lists of column names")
    tolerances = {a: decimal(config.get("tolerances", {}).get(a, "0")) for a in amounts}
    if any(t < 0 for t in tolerances.values()):
        raise ValueError("tolerances cannot be negative")
    inputs, blocked, issues = {}, set(), []
    source_counts = {}
    for side in ("warehouse", "finance"):
        spec = config[side]
        headers, rows = read_source(spec, base)
        indexed, duplicates, found_issues = prepare(headers, rows, spec, keys, amounts, side)
        inputs[side] = indexed
        blocked.update(duplicates)
        issues.extend(found_issues)
        source_counts[side] = len(rows)
        if not rows and not config.get("allow_empty", False):
            issues.append({"source": side, "row": "", "issue": "source is empty"})
    results = []
    for key in sorted(inputs["warehouse"].keys() | inputs["finance"].keys()):
        row = {f"key_{k}": v for k, v in zip(keys, key)}
        left, right = inputs["warehouse"].get(key), inputs["finance"].get(key)
        if key in blocked:
            status = "duplicate_key"
        elif left is None:
            status = "missing_in_warehouse"
        elif right is None:
            status = "missing_in_finance"
        else:
            status = "matched" if all(abs(left[a] - right[a]) <= tolerances[a] for a in amounts) else "amount_mismatch"
        row["status"] = status
        for amount in amounts:
            # Duplicate keys are ambiguous: do not present an arbitrary record as a valid amount.
            row[f"warehouse_{amount}"] = str(left[amount]) if left and key not in blocked else ""
            row[f"finance_{amount}"] = str(right[amount]) if right and key not in blocked else ""
            row[f"difference_{amount}"] = str(left[amount] - right[amount]) if left and right and key not in blocked else ""
        results.append(row)
    totals = {}
    for side in inputs:
        totals[side] = {a: str(sum((v[a] for k, v in inputs[side].items() if k not in blocked), Decimal(0))) for a in amounts}
    counts = dict(Counter(row["status"] for row in results))
    currency_key = config.get("currency_key", "currency" if "currency" in keys else None)
    grouped_totals = {}
    if currency_key:
        if currency_key not in keys:
            raise ValueError("currency_key must be one of the matching keys")
        position = keys.index(currency_key)
        currencies = sorted({k[position] for index in inputs.values() for k in index if k not in blocked})
        for currency in currencies:
            sides = {side: {a: sum((v[a] for k, v in index.items() if k not in blocked and k[position] == currency), Decimal(0)) for a in amounts} for side, index in inputs.items()}
            grouped_totals[currency] = {
                **{side: {a: str(v) for a, v in values.items()} for side, values in sides.items()},
                "difference": {a: str(sides["warehouse"][a] - sides["finance"][a]) for a in amounts},
            }
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "passed": not issues and all(row["status"] == "matched" for row in results),
        "source_rows": source_counts,
        "status_counts": counts,
        "data_quality_issues": len(issues),
        "tolerances": {a: str(t) for a, t in tolerances.items()},
        "totals_for_valid_unique_keys": None if currency_key else totals,
        "total_differences": None if currency_key else {a: str(Decimal(totals["warehouse"][a]) - Decimal(totals["finance"][a])) for a in amounts},
        "totals_by_currency": grouped_totals,
        "totals_complete": not issues,
        "difference_direction": "warehouse minus finance",
    }
    return summary, results, issues


def write_reports(output, summary, results, issues, keys, amounts):
    output.mkdir(parents=True, exist_ok=False)
    (output / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    fields = [f"key_{k}" for k in keys] + ["status"] + [f"{prefix}_{a}" for a in amounts for prefix in ("warehouse", "finance", "difference")]
    for name, rows, columns in (
        ("comparison.csv", results, fields),
        ("exceptions.csv", [r for r in results if r["status"] != "matched"], fields),
        ("data_quality.csv", issues, ["source", "row", "issue"]),
    ):
        with (output / name).open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            for row in rows:
                # Prevent spreadsheet formula execution when Finance opens a report.
                writer.writerow({k: "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")) and not k.startswith(("warehouse_", "finance_", "difference_")) else v for k, v in row.items()})
